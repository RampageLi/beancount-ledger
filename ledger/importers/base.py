"""Common lossless statement model; amounts use Decimal, never binary floats."""
from dataclasses import dataclass, field, asdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import re


def money(value):
    text = str(value).strip().replace(',', '').replace('−', '-')
    text = re.sub(r'[¥￥\s]', '', text)
    if text.startswith('(') and text.endswith(')'):
        text = '-' + text[1:-1]
    result = Decimal(text)
    if not result.is_finite():
        raise ValueError(f'Invalid amount: {value!r}')
    return result


def timestamp(value):
    if isinstance(value, datetime):
        return value
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d',
                '%Y%m%d %H:%M', '%Y%m%d'):
        try:
            return datetime.strptime(str(value).strip(), fmt)
        except ValueError:
            pass
    raise ValueError(f'Invalid date: {value!r}')


@dataclass
class Record:
    source: str
    source_file: str
    source_row: int
    date: datetime
    amount: Decimal  # Original signed amount as shown by the provider.
    currency: str = 'CNY'
    signed_amount: Decimal | None = None  # Account change; unknown for neutral rows.
    payee: str = ''
    narration: str = ''
    direction: str = ''
    status: str = ''
    payment_method: str = ''
    transaction_id: str = ''
    merchant_id: str = ''
    posting_date: str = ''
    balance: Decimal | None = None
    card: str = ''
    page: int | None = None
    raw: dict = field(default_factory=dict)

    def to_dict(self):
        def encode(value):
            if isinstance(value, Decimal):
                return str(value)
            if isinstance(value, datetime):
                return value.isoformat(sep=' ')
            if isinstance(value, dict):
                return {k: encode(v) for k, v in value.items()}
            return value
        return encode(asdict(self))


class Importer:
    source = ''
    account = ''

    def __init__(self, account=None):
        if account:
            self.account = account

    def identify(self, filepath):
        raise NotImplementedError

    def parse_bill(self, filepath):
        raise NotImplementedError

    def extract(self, filepath, existing=None):
        """Review-only Beancount entries. No automatic cross-provider matching."""
        from beancount.core import data
        from beancount.core.amount import Amount
        entries = []
        known = {e.meta.get('import_id') for e in (existing or []) if getattr(e, 'meta', None)}
        for row in self.parse_bill(filepath):
            if row.signed_amount is None:
                continue
            # Only provider transaction IDs are safe deduplication keys.
            identity = f'{row.source}:{row.transaction_id}' if row.transaction_id else None
            if identity and identity in known:
                continue
            meta = data.new_metadata(str(filepath), row.source_row)
            meta.update(source=row.source, payment_method=row.payment_method)
            if identity:
                meta['import_id'] = identity
                known.add(identity)
            entries.append(data.Transaction(
                meta, row.date.date(), '!', row.payee, row.narration,
                frozenset(), frozenset(), [
                    data.Posting(self.account, Amount(row.signed_amount, row.currency), None, None, None, None),
                    data.Posting('Equity:Pending', Amount(-row.signed_amount, row.currency), None, None, None, None),
                ]))
        return entries


def xlsx_rows(filepath):
    from openpyxl import load_workbook
    book = load_workbook(filepath, read_only=True, data_only=True)
    try:
        return list(book.active.iter_rows(values_only=True))
    finally:
        book.close()


def table_rows(rows, required):
    header = None
    for number, row in enumerate(rows, 1):
        values = [str(v).strip() if v is not None else '' for v in row]
        if required.issubset(values):
            header = values
            continue
        if header is not None and any(values):
            yield number, dict(zip(header, row))
    if header is None:
        raise ValueError(f'Missing statement columns: {sorted(required)}')


def clean(row):
    return {k: (str(v).strip() if v is not None else '') for k, v in row.items() if k}
